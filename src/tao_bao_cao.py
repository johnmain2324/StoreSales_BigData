"""Sinh bao cao Word (.docx) tu cac bang/hinh da tao o outputs/tables va
outputs/figures qua cac notebook 01-07. Khong tinh toan lai gi ca - chi doc va
trinh bay lai ket qua da co, nen chay rat nhanh va luon khop voi lan chay
notebook gan nhat.
"""
import datetime

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from cau_hinh import TEN_DE_TAI, PHU_DE, OUT_TABLES, OUT_FIGURES, REPORT, HORIZON, tao_thu_muc

MAU_TIEU_DE = RGBColor(0x0B, 0x0B, 0x0B)


def doc_csv(ten_file, **kwargs):
    return pd.read_csv(OUT_TABLES / ten_file, **kwargs)


def dinh_dang_gia_tri(v):
    if isinstance(v, bool):
        return "Có" if v else "Không"
    if isinstance(v, float):
        return f"{v:,.4f}" if abs(v) < 100 else f"{v:,.2f}"
    return str(v)


def them_bang_tu_df(document, df, so_dong_toi_da=15):
    df_hien_thi = df.head(so_dong_toi_da)
    bang = document.add_table(rows=1, cols=len(df_hien_thi.columns))
    bang.style = "Table Grid"
    hang_tieu_de = bang.rows[0].cells
    for i, ten_cot in enumerate(df_hien_thi.columns):
        hang_tieu_de[i].text = str(ten_cot)
        for p in hang_tieu_de[i].paragraphs:
            for r in p.runs:
                r.bold = True

    for _, hang in df_hien_thi.iterrows():
        cells = bang.add_row().cells
        for i, gia_tri in enumerate(hang):
            cells[i].text = dinh_dang_gia_tri(gia_tri)

    if len(df) > so_dong_toi_da:
        ghi_chu = document.add_paragraph(f"... (hiển thị {so_dong_toi_da}/{len(df)} dòng)")
        ghi_chu.runs[0].italic = True
        ghi_chu.runs[0].font.size = Pt(9)
    document.add_paragraph()


def them_hinh(document, ten_file, chu_thich, do_rong_inch=6.0):
    duong_dan = OUT_FIGURES / ten_file
    if not duong_dan.exists():
        document.add_paragraph(f"[Không tìm thấy hình: {ten_file}]")
        return
    document.add_picture(str(duong_dan), width=Inches(do_rong_inch))
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = document.add_paragraph(chu_thich)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].italic = True
    p.runs[0].font.size = Pt(10)


def them_muc_luc(document):
    p = document.add_paragraph()
    r = p.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-2" \\h \\z \\u'
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_text = OxmlElement("w:t")
    fld_text.text = "Nhấn chuột phải vào mục lục và chọn 'Update Field' để cập nhật."
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    for phan_tu in [fld_begin, instr, fld_sep, fld_text, fld_end]:
        r._r.append(phan_tu)


def them_trang_bia(document):
    document.add_paragraph()
    tieu_de = document.add_paragraph()
    tieu_de.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = tieu_de.add_run(TEN_DE_TAI)
    run.bold = True
    run.font.size = Pt(20)
    run.font.color.rgb = MAU_TIEU_DE

    phu_de = document.add_paragraph()
    phu_de.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = phu_de.add_run(PHU_DE)
    run.font.size = Pt(14)
    run.italic = True

    document.add_paragraph()
    thong_tin = document.add_paragraph()
    thong_tin.alignment = WD_ALIGN_PARAGRAPH.CENTER
    thong_tin.add_run(
        "Nhóm 09 - Mã đề tài 4 - Chuyên đề Xử lý Dữ liệu lớn\n"
        "Trường Đại học Thủ Dầu Một\n\n"
        f"Ngày tạo báo cáo: {datetime.date.today().isoformat()}"
    )
    document.add_page_break()


def them_gioi_thieu(document):
    document.add_heading("1. Giới thiệu", level=1)
    document.add_paragraph(
        "Báo cáo trình bày quy trình xử lý dữ liệu lớn và dự báo chuỗi thời gian trên "
        "bộ dữ liệu Store Sales (Kaggle - Store Sales Time Series Forecasting), gồm "
        "54 cửa hàng và 33 nhóm hàng tại Ecuador. Mục tiêu: dự báo doanh số bán hàng "
        f"cho {HORIZON} ngày tiếp theo (2017-08-16 → 2017-08-31) bằng thuật toán "
        "Random Forest cài đặt trên Apache Spark (chế độ local)."
    )
    document.add_paragraph(
        "Toàn bộ quy trình gồm 7 giai đoạn (notebooks/01 → 07): nạp dữ liệu, cài đặt "
        "MapReduce thuần (không dùng Spark) để minh hoạ nguyên lý, tiền xử lý và hợp "
        "nhất 7 bảng dữ liệu gốc, phân tích khám phá (EDA), xây dựng đặc trưng, huấn "
        "luyện Random Forest, và đánh giá mở rộng (so sánh với Gradient-Boosted Trees)."
    )


def them_moi_truong(document):
    document.add_heading("2. Môi trường và công cụ", level=1)
    document.add_paragraph(
        "Môi trường được kiểm tra tự động bằng src/kiem_tra_moi_truong.py trước khi "
        "bắt đầu xử lý dữ liệu (Giai đoạn 1)."
    )
    df = doc_csv("moi_truong.csv")
    them_bang_tu_df(document, df, so_dong_toi_da=len(df))


def them_du_lieu(document):
    document.add_heading("3. Dữ liệu", level=1)
    document.add_paragraph(
        "Bộ dữ liệu gốc gồm 7 file CSV tải từ Kaggle, được đọc bằng Spark với schema "
        "tường minh và ghi lại dưới dạng Parquet (notebooks/01_du_lieu.ipynb): "
        "train.csv (3.000.888 dòng, 2013-01-01 → 2017-08-15), test.csv (28.512 dòng, "
        "16 ngày cần dự báo), stores.csv (54 cửa hàng), oil.csv (giá dầu WTI theo "
        "ngày), holidays_events.csv (ngày lễ/sự kiện), transactions.csv (số giao dịch "
        "theo cửa hàng/ngày, chỉ có ở giai đoạn train)."
    )


def them_mapreduce(document):
    document.add_heading("4. MapReduce thuần", level=1)
    document.add_paragraph(
        "Để minh hoạ nguyên lý MapReduce trước khi dùng Spark, notebooks/02_mapreduce.ipynb "
        "cài đặt mô hình map → shuffle → reduce bằng Python thuần + multiprocessing "
        "(src/mapreduce/), giải 3 bài toán tổng hợp doanh số theo cửa hàng / nhóm hàng / "
        "tháng trên toàn bộ train.csv (~3 triệu dòng). Kết quả đã được đối chiếu (cross-check) "
        "với Spark groupBy trên cùng dữ liệu và khớp tuyệt đối (sai lệch ~3e-8, chỉ do làm tròn "
        "số thực)."
    )
    document.add_heading("Top 10 nhóm hàng theo tổng doanh số", level=2)
    df = doc_csv("mapreduce_theo_nhom_hang.csv")
    them_bang_tu_df(document, df, so_dong_toi_da=10)
    them_hinh(document, "mapreduce_top10_nhom_hang.png", "Hình 1. Top 10 nhóm hàng theo tổng doanh số (2013-2017).")
    them_hinh(document, "mapreduce_doanh_so_theo_thang.png", "Hình 2. Tổng doanh số theo tháng (toàn bộ chuỗi).")


def them_tien_xu_ly(document):
    document.add_heading("5. Tiền xử lý và hợp nhất dữ liệu", level=1)
    document.add_paragraph(
        "notebooks/03_tien_xu_ly.ipynb hợp nhất train/test với stores, oil, "
        "holidays_events, transactions thành 1 bảng duy nhất theo lưới (store_nbr, date). "
        "Giá dầu được tái lập lịch đầy đủ (1704 ngày, từ 1218 ngày gốc) và nội suy tuyến "
        "tính. Ngày lễ được quy đổi đúng theo 3 phạm vi (National/Regional/Local) và cơ "
        "chế 'transferred' (lễ bị dời ngày), tách riêng các sự kiện đặc biệt (vd. động đất "
        "2016) khỏi cờ ngày lễ thật. Đã kiểm chứng: join không nhân bản dòng (train vẫn "
        "đúng 3.000.888 dòng, test vẫn đúng 28.512 dòng)."
    )


def them_eda(document):
    document.add_heading("6. Phân tích khám phá dữ liệu (EDA)", level=1)
    document.add_paragraph(
        "notebooks/04_eda.ipynb phát hiện một số đặc điểm quan trọng cho bước đặc trưng: "
        "lưới thời gian đầy đủ nhưng ~31% dòng có doanh số = 0, một phần do các cửa hàng "
        "mở muộn (vd. cửa hàng 52 mở 2017-04-20). Kiểm định ADF cho thấy chuỗi tổng doanh "
        "số có xu hướng tăng và khả năng không dừng."
    )
    document.add_heading("Cửa hàng mở muộn (phát hiện từ dữ liệu)", level=2)
    df = doc_csv("eda_cua_hang_mo_muon.csv")
    them_bang_tu_df(document, df, so_dong_toi_da=len(df))
    them_hinh(document, "eda_xu_huong_theo_ngay.png", "Hình 3. Xu hướng tổng doanh số toàn quốc theo ngày.")
    them_hinh(document, "eda_theo_thu.png", "Hình 4. Doanh số trung bình theo thứ trong tuần.")
    them_hinh(document, "eda_gia_dau_vs_doanh_so.png",
              "Hình 5. Giá dầu vs tổng doanh số ngày (lưu ý: có thể là giả tương quan theo thời gian).")
    them_hinh(document, "eda_theo_loai_cua_hang.png", "Hình 6. Doanh số trung bình theo loại cửa hàng.")


def them_dac_trung(document):
    document.add_heading("7. Xây dựng đặc trưng", level=1)
    document.add_paragraph(
        "notebooks/05_dac_trung.ipynb gộp train và test thành 1 khung thời gian liên tục "
        "để tính đặc trưng trễ (lag) và trung bình trượt đúng theo thứ tự thời gian (không "
        "rò rỉ dữ liệu tương lai), cắt bỏ đoạn doanh số = 0 giả trước ngày một nhóm hàng "
        "thực sự được bán tại từng cửa hàng (bỏ ~15,4% dòng). Đặc trưng gồm: thời gian "
        "(năm/tháng/ngày/thứ/tuần/cuối tuần), lag 7/14/28 ngày, trung bình và độ lệch "
        "chuẩn trượt 7/28 ngày, ngày lễ, sự kiện đặc biệt, khuyến mãi, giá dầu."
    )
    document.add_paragraph(
        "Phát hiện quan trọng: đặc trưng cửa sổ ngắn (lag_7, lag_14) bị null dần về cuối "
        f"{HORIZON} ngày dự báo (vì 'hôm qua' cũng là ngày chưa biết trước), trong khi "
        "lag_28/trung_bình_trượt_28 không bao giờ null trên test - nên được ưu tiên cho "
        "dự báo nhiều bước."
    )


def them_mo_hinh(document):
    document.add_heading("8. Mô hình Random Forest", level=1)
    document.add_paragraph(
        "notebooks/06_random_forest.ipynb huấn luyện RandomForestRegressor (Spark MLlib, "
        "60 cây, độ sâu tối đa 12) dự báo log1p(doanh số). Chia tập theo thời gian (không "
        f"random split): giữ lại {HORIZON} ngày cuối của train làm kiểm định, mô phỏng "
        "đúng tình huống dự báo thực tế."
    )
    document.add_heading("Kết quả trên tập kiểm định", level=2)
    df = doc_csv("danh_gia_random_forest.csv")
    them_bang_tu_df(document, df, so_dong_toi_da=len(df))
    document.add_paragraph(
        "Random Forest đạt RMSLE thấp hơn đáng kể so với baseline mùa vụ ngây thơ "
        "(dự báo bằng doanh số cùng thứ 4 tuần trước), chứng minh mô hình học được "
        "quy luật thực sự ngoài việc lặp lại mùa vụ."
    )
    them_hinh(document, "random_forest_do_quan_trong.png", "Hình 7. Top 12 đặc trưng quan trọng nhất.")


def them_danh_gia_mo_rong(document):
    document.add_heading("9. Đánh giá mở rộng", level=1)
    document.add_paragraph(
        "notebooks/07_danh_gia_mo_rong.ipynb phân tích sâu lỗi dự báo và so sánh với "
        "Gradient-Boosted Trees (GBT) trên cùng đặc trưng/tập dữ liệu."
    )
    document.add_heading("So sánh Random Forest vs GBT", level=2)
    df = doc_csv("so_sanh_random_forest_gbt.csv")
    them_bang_tu_df(document, df, so_dong_toi_da=len(df))

    document.add_heading("Lỗi theo nhóm hàng (5 cao nhất / 5 thấp nhất)", level=2)
    df_family = doc_csv("danh_gia_rmsle_theo_family.csv")
    them_bang_tu_df(document, pd.concat([df_family.head(5), df_family.tail(5)]), so_dong_toi_da=10)
    them_hinh(document, "danh_gia_rmsle_theo_family.png", "Hình 8. RMSLE theo nhóm hàng.")

    document.add_heading("Lỗi theo bước dự báo và theo loại cửa hàng", level=2)
    them_hinh(document, "danh_gia_rmsle_theo_buoc.png", "Hình 9. RMSLE theo bước dự báo trong horizon 16 ngày.")
    them_hinh(document, "danh_gia_rmsle_theo_loai.png", "Hình 10. RMSLE theo loại cửa hàng.")
    them_hinh(document, "danh_gia_thuc_te_vs_du_bao.png",
              "Hình 11. Doanh số thực tế vs dự báo (mẫu 2% tập kiểm định).")


def them_ket_luan(document):
    document.add_heading("10. Kết luận và hướng phát triển", level=1)
    document.add_paragraph(
        "Quy trình đã hoàn thành đầy đủ từ nạp dữ liệu, MapReduce thuần, tiền xử lý, EDA, "
        "xây dựng đặc trưng, đến huấn luyện và đánh giá mô hình Random Forest trên Apache "
        "Spark. Mô hình vượt trội baseline mùa vụ và nhỉnh hơn GBT trong cùng điều kiện."
    )
    document.add_paragraph("Hướng phát triển tiếp theo:")
    for muc in [
        "Dò siêu tham số (numTrees, maxDepth) bằng TrainValidationSplit theo phân chia "
        "thời gian (không random k-fold để tránh rò rỉ dữ liệu tương lai).",
        "Huấn luyện mô hình riêng cho các nhóm hàng có RMSLE cao (vd. SCHOOL AND OFFICE "
        "SUPPLIES, LINGERIE) thay vì dùng chung 1 mô hình cho tất cả.",
        "Thử nghiệm mô hình ensemble kết hợp Random Forest và GBT.",
        "Bổ sung đặc trưng từ sự kiện đặc biệt (động đất 2016) và tác động trễ của "
        "khuyến mãi (lag onpromotion).",
    ]:
        document.add_paragraph(muc, style="List Bullet")


def main():
    tao_thu_muc()
    document = Document()

    them_trang_bia(document)
    document.add_heading("Mục lục", level=1)
    them_muc_luc(document)
    document.add_page_break()

    them_gioi_thieu(document)
    them_moi_truong(document)
    them_du_lieu(document)
    them_mapreduce(document)
    them_tien_xu_ly(document)
    them_eda(document)
    them_dac_trung(document)
    them_mo_hinh(document)
    them_danh_gia_mo_rong(document)
    them_ket_luan(document)

    duong_dan_ra = REPORT / "BaoCao_Nhom09_DeTai4_StoreSales.docx"
    document.save(str(duong_dan_ra))
    print("Da luu bao cao:", duong_dan_ra)


if __name__ == "__main__":
    main()
